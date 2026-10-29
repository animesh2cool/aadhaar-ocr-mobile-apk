"""
AadhaarScan Mobile — Improved Kivy + Google ML Kit edition

Main improvements:
- OCR runs on a background thread; Kivy widgets are updated on the main thread.
- Prevents multiple simultaneous scans.
- Better image-load and OCR error handling.
- Fixes Excel export alignment bug and handles save errors.
- Validates dates and normalizes common Aadhaar-number OCR spacing.
- Masks the Aadhaar number in the UI by default (tap Show/Hide).
- Avoids requesting broad legacy storage permissions at startup.
- Stops the camera when the camera popup closes.
- Stores temporary camera captures in the app's private storage.

Important:
- Android ML Kit requires the appropriate Buildozer/Python-for-Android Gradle dependency.
- This file alone cannot configure or guarantee the APK's native dependencies.
- The Kivy FileChooser is retained as a basic picker. On Android it may only show
  files visible to the app. A native Android Photo Picker/SAF integration is a
  separate improvement if you need to select directly from Gallery/Downloads.
- OCR output is not verified identity data. Always review it before exporting.
"""

import os
import re
import threading
import datetime
import traceback
from typing import Dict, Any, List, Optional

from kivy.utils import platform

# Desktop aspect-ratio simulation. Set before importing/creating the Window.
if platform != "android":
    from kivy.config import Config
    Config.set("graphics", "width", "400")
    Config.set("graphics", "height", "800")
    Config.set("graphics", "resizable", "1")

from kivy.app import App
from kivy.core.window import Window
from kivy.metrics import dp
from kivy.graphics import Color, RoundedRectangle, Rectangle
from kivy.uix.boxlayout import BoxLayout
from kivy.uix.gridlayout import GridLayout
from kivy.uix.scrollview import ScrollView
from kivy.uix.label import Label
from kivy.uix.textinput import TextInput
from kivy.uix.button import Button
from kivy.uix.image import Image
from kivy.uix.camera import Camera
from kivy.uix.popup import Popup
from kivy.clock import Clock

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

Window.softinput_mode = "below_target"
Window.clearcolor = (0.05, 0.07, 0.12, 1.0)


class StyledCard(BoxLayout):
    """A rounded background card."""
    def __init__(self, bg_color=(0.10, 0.13, 0.22, 1.0), radius=14, **kwargs):
        super().__init__(**kwargs)
        self.bg_color = bg_color
        self.radius = radius
        with self.canvas.before:
            Color(*self.bg_color)
            self.rect = RoundedRectangle(
                size=self.size, pos=self.pos, radius=[dp(self.radius)]
            )
        self.bind(size=self._update_rect, pos=self._update_rect)

    def _update_rect(self, *_args):
        self.rect.size = self.size
        self.rect.pos = self.pos


class AadhaarMobileApp(App):
    def build(self):
        self.title = "AadhaarScan"
        self.current_image_path: Optional[str] = None
        self.camera_widget: Optional[Camera] = None
        self.camera_popup: Optional[Popup] = None
        self._ocr_running = False
        self._scan_lock = threading.Lock()
        self._scan_id = 0
        self._last_aadhaar_value = ""
        self._aadhaar_visible = False

        root = BoxLayout(
            orientation="vertical", padding=dp(12), spacing=dp(8)
        )
        with root.canvas.before:
            Color(0.05, 0.07, 0.12, 1.0)
            self.bg_rect = Rectangle(size=root.size, pos=root.pos)
        root.bind(size=self._update_bg, pos=self._update_bg)

        # Header
        header = BoxLayout(size_hint_y=None, height=dp(46), spacing=dp(8))
        title = Label(
            text="[b]AadhaarScan[/b] [color=818cf8]ML Kit[/color]",
            markup=True, font_size="18sp",
            halign="left", valign="middle"
        )
        title.bind(size=title.setter("text_size"))
        header.add_widget(title)

        sample_btn = self._button(
            "Sample", (0.38, 0.40, 0.95, 0.9), width=dp(88)
        )
        sample_btn.bind(on_press=self.load_sample)
        header.add_widget(sample_btn)
        root.add_widget(header)

        # Scrollable content
        scroll = ScrollView(size_hint=(1, 1), do_scroll_x=False)
        content = BoxLayout(
            orientation="vertical", spacing=dp(10), size_hint_y=None
        )
        content.bind(minimum_height=content.setter("height"))

        preview_card = StyledCard(
            bg_color=(0.08, 0.11, 0.18, 1.0),
            orientation="vertical", padding=dp(6),
            size_hint_y=None, height=dp(210)
        )
        self.img_preview = Image(
            source="", allow_stretch=True, keep_ratio=True
        )
        preview_card.add_widget(self.img_preview)
        content.add_widget(preview_card)

        self.status_label = Label(
            text="Scan a document, select an image, or load a sample.",
            font_size="12sp", color=(0.65, 0.75, 0.85, 1),
            size_hint_y=None, height=dp(36),
            halign="center", valign="middle"
        )
        self.status_label.bind(size=self.status_label.setter("text_size"))
        content.add_widget(self.status_label)

        form_card = StyledCard(
            bg_color=(0.09, 0.12, 0.20, 1.0),
            orientation="vertical", padding=dp(12), spacing=dp(6),
            size_hint_y=None
        )
        form_card.bind(minimum_height=form_card.setter("height"))

        self.input_name = self._add_field(
            form_card, "FULL NAME", "Enter full name"
        )

        # Aadhaar field + show/hide control
        form_card.add_widget(self._field_label("AADHAAR NUMBER"))
        aadhaar_row = BoxLayout(
            orientation="horizontal", spacing=dp(6),
            size_hint_y=None, height=dp(42)
        )
        self.input_aadhaar = TextInput(
            hint_text="XXXX XXXX XXXX", multiline=False,
            password=True, size_hint=(1, 1), font_size="16sp",
            background_color=(0.18, 0.12, 0.15, 1),
            foreground_color=(1, 0.65, 0.65, 1),
            cursor_color=(1, 1, 1, 1)
        )
        self.input_aadhaar.bind(text=self._on_aadhaar_edited)
        aadhaar_row.add_widget(self.input_aadhaar)
        self.btn_toggle_aadhaar = self._button(
            "Show", (0.25, 0.30, 0.45, 1), width=dp(70)
        )
        self.btn_toggle_aadhaar.bind(on_press=self.toggle_aadhaar_visibility)
        aadhaar_row.add_widget(self.btn_toggle_aadhaar)
        form_card.add_widget(aadhaar_row)

        row_cols = GridLayout(
            cols=2, spacing=dp(8), size_hint_y=None, height=dp(64)
        )
        dob_col = BoxLayout(orientation="vertical", spacing=dp(2))
        dob_col.add_widget(self._field_label("DATE OF BIRTH"))
        self.input_dob = TextInput(
            hint_text="DD/MM/YYYY or YYYY", multiline=False,
            size_hint_y=None, height=dp(40), font_size="13sp",
            background_color=(0.14, 0.18, 0.28, 1),
            foreground_color=(1, 1, 1, 1)
        )
        dob_col.add_widget(self.input_dob)
        row_cols.add_widget(dob_col)

        gender_col = BoxLayout(orientation="vertical", spacing=dp(2))
        gender_col.add_widget(self._field_label("GENDER"))
        self.input_gender = TextInput(
            hint_text="MALE / FEMALE", multiline=False,
            size_hint_y=None, height=dp(40), font_size="13sp",
            background_color=(0.14, 0.18, 0.28, 1),
            foreground_color=(1, 1, 1, 1)
        )
        gender_col.add_widget(self.input_gender)
        row_cols.add_widget(gender_col)
        form_card.add_widget(row_cols)
        content.add_widget(form_card)

        privacy_note = Label(
            text="Review every extracted field before saving. Keep identity data private.",
            font_size="11sp", color=(0.72, 0.76, 0.84, 1),
            size_hint_y=None, height=dp(34),
            halign="left", valign="middle"
        )
        privacy_note.bind(size=privacy_note.setter("text_size"))
        content.add_widget(privacy_note)

        scroll.add_widget(content)
        root.add_widget(scroll)

        # Actions remain fixed at the bottom
        actions = BoxLayout(
            orientation="vertical", spacing=dp(8),
            size_hint_y=None, height=dp(102)
        )
        row_inputs = GridLayout(
            cols=2, spacing=dp(8), size_hint_y=None, height=dp(44)
        )
        btn_cam = self._button("Live Camera", (0.38, 0.40, 0.95, 1))
        btn_cam.bind(on_press=self.open_camera_view)
        row_inputs.add_widget(btn_cam)

        btn_files = self._button("Select Image", (0.25, 0.30, 0.45, 1))
        btn_files.bind(on_press=self.open_file_selector)
        row_inputs.add_widget(btn_files)
        actions.add_widget(row_inputs)

        self.btn_export = self._button(
            "Export to Excel (.xlsx)", (0.06, 0.62, 0.43, 1),
            height=dp(46)
        )
        self.btn_export.bind(on_press=self.export_to_excel)
        actions.add_widget(self.btn_export)
        root.add_widget(actions)
        return root

    def _button(self, text, color, width=None, height=None):
        kwargs = {
            "text": text,
            "background_normal": "",
            "background_color": color,
            "font_size": "13sp",
            "bold": True,
        }
        if width is not None:
            kwargs["size_hint_x"] = None
            kwargs["width"] = width
        if height is not None:
            kwargs["size_hint_y"] = None
            kwargs["height"] = height
        return Button(**kwargs)

    def _field_label(self, text):
        return Label(
            text=text, font_size="11sp",
            color=(0.6, 0.7, 0.85, 1),
            size_hint_y=None, height=dp(16),
            halign="left", valign="middle"
        )

    def _add_field(self, parent, label, hint):
        parent.add_widget(self._field_label(label))
        field = TextInput(
            hint_text=hint, multiline=False,
            size_hint_y=None, height=dp(40), font_size="14sp",
            background_color=(0.14, 0.18, 0.28, 1),
            foreground_color=(1, 1, 1, 1),
            cursor_color=(1, 1, 1, 1)
        )
        parent.add_widget(field)
        return field

    def on_start(self):
        # Do not request broad storage permissions at startup.
        # Camera permission is requested only when the camera is opened.
        pass

    def on_pause(self):
        self._stop_camera()
        return True

    def on_resume(self):
        return None

    def on_stop(self):
        self._stop_camera()

    def _update_bg(self, instance, _value):
        self.bg_rect.size = instance.size
        self.bg_rect.pos = instance.pos

    def get_app_storage_dir(self):
        """Return a writable app-specific directory."""
        if platform == "android":
            try:
                from android.storage import app_storage_path
                path = app_storage_path()
                if path:
                    os.makedirs(path, exist_ok=True)
                    return path
            except Exception:
                pass
        path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "app_data")
        os.makedirs(path, exist_ok=True)
        return path

    def _request_camera_permission(self, on_granted):
        if platform != "android":
            on_granted()
            return
        try:
            from android.permissions import request_permissions, Permission, check_permission
            if check_permission(Permission.CAMERA):
                on_granted()
                return

            def permission_callback(_permissions, grant_results):
                granted = bool(grant_results) and all(bool(v) for v in grant_results)
                if granted:
                    Clock.schedule_once(lambda _dt: on_granted(), 0)
                else:
                    Clock.schedule_once(
                        lambda _dt: self.show_popup(
                            "Camera permission",
                            "Camera permission was denied. You can still select an image."
                        ), 0
                    )

            request_permissions([Permission.CAMERA], permission_callback)
        except Exception as exc:
            self.show_popup(
                "Camera unavailable",
                "Could not request camera permission. Check Android permissions.\n"
                + str(exc)
            )

    def open_camera_view(self, _instance=None):
        if self._ocr_running:
            self.show_popup("Please wait", "Wait for the current scan to finish.")
            return
        self._request_camera_permission(self._create_camera_popup)

    def _create_camera_popup(self):
        cam_layout = BoxLayout(
            orientation="vertical", padding=dp(10), spacing=dp(10)
        )
        try:
            # Resolution is a request; the camera provider/device may choose another.
            self.camera_widget = Camera(
                play=True, resolution=(1280, 720), index=0
            )
            cam_layout.add_widget(self.camera_widget)
        except Exception as exc:
            self.camera_widget = None
            self.show_popup("Camera error", f"Could not initialize the camera.\n{exc}")
            return

        btn_row = BoxLayout(
            size_hint_y=None, height=dp(48), spacing=dp(10)
        )
        btn_capture = self._button(
            "Capture & Scan", (0.38, 0.40, 0.95, 1)
        )
        btn_close = self._button("Close", (0.3, 0.3, 0.3, 1))
        btn_row.add_widget(btn_capture)
        btn_row.add_widget(btn_close)
        cam_layout.add_widget(btn_row)

        popup = Popup(
            title="Align the document and capture",
            content=cam_layout, size_hint=(0.96, 0.88),
            auto_dismiss=False
        )
        self.camera_popup = popup

        def close_camera(*_args):
            self._stop_camera()
            if popup.parent:
                popup.dismiss()

        def capture(_btn):
            camera = self.camera_widget
            if camera is None or camera.texture is None:
                self.show_popup(
                    "Camera not ready",
                    "Wait until the camera preview appears, then try again."
                )
                return
            try:
                filename = "scan_{}.png".format(
                    datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
                )
                path = os.path.join(self.get_app_storage_dir(), filename)
                camera.export_to_png(path)
                if not os.path.isfile(path) or os.path.getsize(path) == 0:
                    raise IOError("The camera did not produce an image.")
                close_camera()
                self.process_image_path(path)
            except Exception as exc:
                self.status_label.text = "Capture failed."
                self.show_popup("Capture error", str(exc))

        btn_capture.bind(on_press=capture)
        btn_close.bind(on_press=close_camera)
        popup.bind(on_dismiss=lambda *_args: self._stop_camera())
        popup.open()

    def _stop_camera(self):
        try:
            if self.camera_widget is not None:
                self.camera_widget.play = False
        except Exception:
            pass
        self.camera_widget = None
        self.camera_popup = None

    def open_file_selector(self, _instance=None):
        if self._ocr_running:
            self.show_popup("Please wait", "Wait for the current scan to finish.")
            return
        try:
            from kivy.uix.filechooser import FileChooserListView
            chooser = FileChooserListView(
                path=self.get_app_storage_dir(),
                filters=["*.png", "*.jpg", "*.jpeg", "*.webp", "*.bmp"]
            )
        except Exception as exc:
            self.show_popup("File picker error", str(exc))
            return

        popup_layout = BoxLayout(
            orientation="vertical", padding=dp(10), spacing=dp(10)
        )
        popup_layout.add_widget(chooser)
        btn_row = BoxLayout(
            size_hint_y=None, height=dp(44), spacing=dp(10)
        )
        btn_select = self._button("Select", (0.38, 0.40, 0.95, 1))
        btn_cancel = self._button("Cancel", (0.3, 0.3, 0.3, 1))
        btn_row.add_widget(btn_select)
        btn_row.add_widget(btn_cancel)
        popup_layout.add_widget(btn_row)
        popup = Popup(
            title="Select an image",
            content=popup_layout, size_hint=(0.96, 0.88)
        )

        def select_file(_btn):
            if not chooser.selection:
                self.show_popup("Select image", "Choose an image first.")
                return
            selected_path = chooser.selection[0]
            if not os.path.isfile(selected_path):
                self.show_popup(
                    "Image unavailable",
                    "This picker can only open a local file path. "
                    "On some Android versions, use an app-private copy or add native Photo Picker support."
                )
                return
            popup.dismiss()
            self.process_image_path(selected_path)

        btn_select.bind(on_press=select_file)
        btn_cancel.bind(on_press=popup.dismiss)
        popup.open()

    def load_sample(self, _instance=None):
        base_dir = os.path.dirname(os.path.abspath(__file__))
        candidates = [
            os.path.join(base_dir, "sample_aadhaar.png"),
            os.path.join(base_dir, "Screenshot 2026-10-28 133028.png"),
        ]
        sample_path = next((p for p in candidates if os.path.isfile(p)), None)
        if not sample_path:
            self.show_popup(
                "Sample image not found",
                "Place sample_aadhaar.png beside this Python file, or select a real image."
            )
            return
        self.process_image_path(sample_path)

    def process_image_path(self, path: str):
        if self._ocr_running:
            self.show_popup("Please wait", "A scan is already running.")
            return
        if not path or not os.path.isfile(path):
            self.show_popup("Image error", "The selected image could not be found.")
            return

        extension = os.path.splitext(path)[1].lower()
        if extension not in {".png", ".jpg", ".jpeg", ".webp", ".bmp"}:
            self.show_popup("Unsupported file", "Choose a PNG, JPG, JPEG, WEBP, or BMP image.")
            return

        try:
            if os.path.getsize(path) > 40 * 1024 * 1024:
                self.show_popup(
                    "Image too large",
                    "This image is larger than 40 MB. Choose a smaller image."
                )
                return
        except OSError:
            self.show_popup("Image error", "Could not read the selected image.")
            return

        self._ocr_running = True
        self._scan_id += 1
        scan_id = self._scan_id
        self.current_image_path = path
        self.img_preview.source = path
        self.img_preview.reload()
        self.status_label.text = "Processing image… please wait."
        self.btn_export.disabled = True

        worker = threading.Thread(
            target=self._ocr_worker,
            args=(path, scan_id),
            daemon=True
        )
        worker.start()

    def _ocr_worker(self, path: str, scan_id: int):
        """Background worker. Do not touch Kivy widgets in this method."""
        lines: List[str] = []
        error_message = ""
        try:
            if platform == "android":
                lines = self._run_mlkit_ocr(path)
            else:
                try:
                    from aadhaar_ocr import AadhaarOCR
                    ocr = AadhaarOCR(
                        languages=["en", "hi"], use_gpu=False, verbose=False
                    )
                    extracted = ocr.extract_text(path)
                    if extracted:
                        lines = [
                            str(line).strip()
                            for line in extracted
                            if str(line).strip()
                        ]
                except Exception as exc:
                    error_message = (
                        "Desktop OCR is unavailable. Install/configure "
                        "aadhaar_ocr and its OCR dependencies. Details: "
                        + str(exc)
                    )
        except Exception as exc:
            # Keep personal document contents out of logs.
            error_message = "OCR failed: " + str(exc)

        parsed = self.parse_aadhaar_text(lines)
        Clock.schedule_once(
            lambda _dt: self._finish_ocr(scan_id, parsed, error_message), 0
        )

    def _run_mlkit_ocr(self, path: str) -> List[str]:
        """Run ML Kit synchronously inside the worker thread, not the UI thread."""
        from jnius import autoclass

        BitmapFactory = autoclass("android.graphics.BitmapFactory")
        InputImage = autoclass("com.google.mlkit.vision.common.InputImage")
        TextRecognition = autoclass("com.google.mlkit.vision.text.TextRecognition")
        Options = autoclass(
            "com.google.mlkit.vision.text.latin.TextRecognizerOptions"
        )
        Tasks = autoclass("com.google.android.gms.tasks.Tasks")

        bitmap = BitmapFactory.decodeFile(path)
        if bitmap is None:
            raise ValueError("Android could not decode this image.")

        recognizer = None
        try:
            input_image = InputImage.fromBitmap(bitmap, 0)
            recognizer = TextRecognition.getClient(Options.DEFAULT_OPTIONS)
            await_method = getattr(Tasks, "await")
            vision_text = await_method(recognizer.process(input_image))
            text = str(vision_text.getText())
            return [line.strip() for line in text.splitlines() if line.strip()]
        finally:
            if recognizer is not None:
                try:
                    recognizer.close()
                except Exception:
                    pass
            try:
                bitmap.recycle()
            except Exception:
                pass

    def _finish_ocr(self, scan_id: int, parsed: Dict[str, Any], error_message: str):
        """Runs on Kivy's main thread."""
        if scan_id != self._scan_id:
            return
        self._ocr_running = False
        self.btn_export.disabled = False

        if error_message and not any(parsed.values()):
            self.status_label.text = "Could not extract text. You can enter details manually."
            self.show_popup("OCR problem", error_message)
            return

        self.input_name.text = parsed.get("name", "")
        self.input_aadhaar.text = parsed.get("aadhaar_number", "")
        self.input_dob.text = parsed.get("dob", "")
        self.input_gender.text = parsed.get("gender", "")

        if any(parsed.values()):
            self.status_label.text = "OCR finished. Review and correct every field."
        elif error_message:
            self.status_label.text = "Image loaded, but OCR was unavailable."
        else:
            self.status_label.text = "No reliable fields found. Enter or correct them manually."

    def parse_aadhaar_text(self, lines: List[str]) -> Dict[str, Any]:
        full_text = " \n ".join(lines)
        data: Dict[str, Any] = {
            "aadhaar_number": "", "name": "", "dob": "", "gender": ""
        }

        # Find candidate 12-digit numbers with flexible whitespace.
        digit_candidates = re.findall(
            r"(?<!\d)(?:\d[\s\-]?){11}\d(?!\d)", full_text
        )
        if digit_candidates:
            for candidate in digit_candidates:
                digits = re.sub(r"\D", "", candidate)
                if len(digits) == 12:
                    data["aadhaar_number"] = (
                        f"{digits[:4]} {digits[4:8]} {digits[8:]}"
                    )
                    break

        # Prefer a labeled full DOB; accept year of birth as a year only.
        dob_patterns = [
            r"(?:DOB|D\.O\.B|Date\s*of\s*Birth|जन्म\s*तिथि|जन्म\s*तारीख)"
            r"[\s:\-]*(\d{2}[/\-\.]\d{2}[/\-\.]\d{4})",
            r"\b(\d{2}[/\-\.]\d{2}[/\-\.]\d{4})\b",
        ]
        for pattern in dob_patterns:
            match = re.search(pattern, full_text, re.IGNORECASE)
            if match:
                raw = match.group(1).replace("-", "/").replace(".", "/")
                try:
                    datetime.datetime.strptime(raw, "%d/%m/%Y")
                    data["dob"] = raw
                except ValueError:
                    pass
                if data["dob"]:
                    break
        if not data["dob"]:
            year_match = re.search(
                r"(?:Year\s*of\s*Birth|जन्म\s*वर्ष)[\s:\-]*(\d{4})\b",
                full_text, re.IGNORECASE
            )
            if year_match:
                year = int(year_match.group(1))
                if 1900 <= year <= datetime.datetime.now().year:
                    data["dob"] = str(year)

        # Check FEMALE before MALE to avoid partial-word confusion.
        if re.search(r"\b(FEMALE|महिला)\b", full_text, re.IGNORECASE):
            data["gender"] = "FEMALE"
        elif re.search(r"\b(MALE|पुरुष)\b", full_text, re.IGNORECASE):
            data["gender"] = "MALE"
        elif re.search(r"\b(TRANSGENDER|उभयलिंगी)\b", full_text, re.IGNORECASE):
            data["gender"] = "TRANSGENDER"

        # Conservative English-name heuristic. User should always verify it.
        ignore = (
            "government", "india", "bharat", "sarkar", "unique",
            "authority", "dob", "male", "female", "aadhaar", "aadhar",
            "card", "enrollment", "enrolment", "year of birth"
        )
        for line in lines:
            candidate = line.strip()
            lower = candidate.lower()
            if any(word in lower for word in ignore):
                continue
            if re.search(r"\d", candidate):
                continue
            if re.fullmatch(r"[A-Za-z][A-Za-z .'-]{2,39}", candidate):
                words = candidate.split()
                if 1 <= len(words) <= 5:
                    data["name"] = candidate
                    break

        return data

    def _on_aadhaar_edited(self, _instance, value):
        # Keep the current value in memory for the show/hide control.
        self._last_aadhaar_value = value

    def toggle_aadhaar_visibility(self, _instance=None):
        self._aadhaar_visible = not self._aadhaar_visible
        self.input_aadhaar.password = not self._aadhaar_visible
        self.btn_toggle_aadhaar.text = "Hide" if self._aadhaar_visible else "Show"

    def _validate_form(self):
        name = self.input_name.text.strip()
        aadhaar_raw = self.input_aadhaar.text.strip()
        dob = self.input_dob.text.strip()
        gender = self.input_gender.text.strip().upper()

        digits = re.sub(r"\D", "", aadhaar_raw)
        if aadhaar_raw and len(digits) != 12:
            return None, "Aadhaar number should contain exactly 12 digits."
        if dob:
            if re.fullmatch(r"\d{4}", dob):
                year = int(dob)
                if not (1900 <= year <= datetime.datetime.now().year):
                    return None, "Please check the year of birth."
            else:
                normalized = dob.replace("-", "/").replace(".", "/")
                try:
                    datetime.datetime.strptime(normalized, "%d/%m/%Y")
                except ValueError:
                    return None, "Date of birth must be a valid DD/MM/YYYY date or a four-digit year."
                dob = normalized
        if gender and gender not in {"MALE", "FEMALE", "TRANSGENDER", "OTHER"}:
            return None, "Gender should be MALE, FEMALE, TRANSGENDER, or OTHER."

        aadhaar_formatted = (
            f"{digits[:4]} {digits[4:8]} {digits[8:]}" if digits else ""
        )
        if not name and not aadhaar_formatted:
            return None, "Please scan a document or enter a name or Aadhaar number."
        return {
            "name": name,
            "aadhaar_number": aadhaar_formatted,
            "dob": dob,
            "gender": gender,
        }, ""

    def export_to_excel(self, _instance=None):
        record, validation_error = self._validate_form()
        if validation_error:
            self.show_popup("Check details", validation_error)
            return

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Aadhaar Card"
            headers = [
                "S.No.", "Aadhaar Number", "Full Name",
                "Date of Birth", "Gender", "Export Date"
            ]
            ws.append(headers)
            ws.append([
                1, record["aadhaar_number"], record["name"],
                record["dob"], record["gender"],
                datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            ])

            header_fill = PatternFill(
                start_color="1E293B", end_color="1E293B", fill_type="solid"
            )
            header_font = Font(
                name="Arial", size=11, bold=True, color="FFFFFF"
            )
            side = Side(border_style="thin", color="CBD5E1")
            border = Border(top=side, bottom=side, left=side, right=side)
            center = Alignment(horizontal="center", vertical="center")
            left = Alignment(horizontal="left", vertical="center")

            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center
                cell.border = border

            for cell in ws[2]:
                cell.border = border
                cell.alignment = left if cell.column == 3 else center
                # Keep identifier values as text, preserving leading zeroes.
                if cell.column == 2:
                    cell.number_format = "@"

            ws.freeze_panes = "A2"
            ws.auto_filter.ref = f"A1:F2"
            ws.row_dimensions[1].height = 24
            for column_cells in ws.columns:
                max_length = max(
                    len(str(cell.value or "")) for cell in column_cells
                )
                ws.column_dimensions[
                    column_cells[0].column_letter
                ].width = min(max(max_length + 4, 12), 34)

            filename = "Aadhaar_Data_{}.xlsx".format(
                datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
            )
            export_dir = self.get_app_storage_dir()
            export_path = os.path.join(export_dir, filename)
            wb.save(export_path)
            wb.close()

            self.show_popup(
                "Export successful",
                "Excel file saved in the app's storage:\n\n"
                + export_path
                + "\n\nUse your file manager or add Android file-sharing integration to move/share it."
            )
        except Exception as exc:
            self.show_popup(
                "Export failed",
                "Could not save the Excel file. Check available storage and try again.\n\n"
                + str(exc)
            )

    def show_popup(self, title: str, message: str):
        layout = BoxLayout(
            orientation="vertical", padding=dp(14), spacing=dp(10)
        )
        label = Label(
            text=message, font_size="13sp",
            halign="center", valign="middle"
        )
        label.bind(size=label.setter("text_size"))
        layout.add_widget(label)
        btn_close = self._button(
            "OK", (0.38, 0.40, 0.95, 1), height=dp(42)
        )
        layout.add_widget(btn_close)
        popup = Popup(
            title=title, content=layout, size_hint=(0.90, 0.48)
        )
        btn_close.bind(on_press=popup.dismiss)
        popup.open()


if __name__ == "__main__":
    AadhaarMobileApp().run()
