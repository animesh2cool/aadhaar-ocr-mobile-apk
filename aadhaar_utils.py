"""
aadhaar_utils.py - pure-Python helpers for AadhaarScan.

No Kivy / Android imports in here, so everything can be unit-tested on a laptop
and the app starts faster (openpyxl / Pillow are imported lazily).
"""

import datetime
import os
import re
from collections import Counter
from typing import Any, Dict, List

# --------------------------------------------------------------------------
# Verhoeff checksum (the last digit of every real Aadhaar number)
# --------------------------------------------------------------------------
_D = [
    [0, 1, 2, 3, 4, 5, 6, 7, 8, 9],
    [1, 2, 3, 4, 0, 6, 7, 8, 9, 5],
    [2, 3, 4, 0, 1, 7, 8, 9, 5, 6],
    [3, 4, 0, 1, 2, 8, 9, 5, 6, 7],
    [4, 0, 1, 2, 3, 9, 5, 6, 7, 8],
    [5, 9, 8, 7, 6, 0, 4, 3, 2, 1],
    [6, 5, 9, 8, 7, 1, 0, 4, 3, 2],
    [7, 6, 5, 9, 8, 2, 1, 0, 4, 3],
    [8, 7, 6, 5, 9, 3, 2, 1, 0, 4],
    [9, 8, 7, 6, 5, 4, 3, 2, 1, 0],
]
_P = [
    [0, 1, 2, 3, 4, 5, 6, 7, 8, 9],
    [1, 5, 7, 6, 2, 8, 3, 0, 9, 4],
    [5, 8, 0, 3, 7, 9, 6, 1, 4, 2],
    [8, 9, 1, 6, 0, 4, 3, 5, 2, 7],
    [9, 4, 5, 3, 1, 2, 6, 8, 7, 0],
    [4, 2, 8, 6, 5, 7, 3, 9, 0, 1],
    [2, 7, 9, 3, 8, 0, 6, 4, 1, 5],
    [7, 0, 4, 6, 9, 1, 3, 2, 5, 8],
]
_INV = [0, 4, 3, 2, 1, 5, 6, 7, 8, 9]


def digits_only(text: str) -> str:
    return re.sub(r"\D", "", text or "")


def format_aadhaar(value: str) -> str:
    d = digits_only(value)[:12]
    return " ".join(d[i:i + 4] for i in range(0, len(d), 4))


def mask_aadhaar(value: str) -> str:
    """'1234 5678 9012' -> 'XXXX XXXX 9012'."""
    d = digits_only(value)
    if not d:
        return ""
    if len(d) <= 4:
        return "X" * len(d)
    masked = "X" * (len(d) - 4) + d[-4:]
    return " ".join(masked[i:i + 4] for i in range(0, len(masked), 4))


def verhoeff_valid(number: str) -> bool:
    if not number or not number.isdigit():
        return False
    c = 0
    for i, ch in enumerate(reversed(number)):
        c = _D[c][_P[i % 8][int(ch)]]
    return c == 0


def verhoeff_check_digit(number: str) -> str:
    """Check digit to append to `number` (used by the tests)."""
    c = 0
    for i, ch in enumerate(reversed(number)):
        c = _D[c][_P[(i + 1) % 8][int(ch)]]
    return str(_INV[c])


def is_valid_aadhaar(value: str) -> bool:
    d = digits_only(value)
    return len(d) == 12 and d[0] not in "01" and verhoeff_valid(d)


# --------------------------------------------------------------------------
# Text parsing
# --------------------------------------------------------------------------
# OCR often reads digits as look-alike letters inside the number line.
_OCR_FIX = str.maketrans({"O": "0", "o": "0", "I": "1", "l": "1", "|": "1",
                          "S": "5", "B": "8", "Z": "2"})
_NUM_CH = r"[0-9OoIl|SBZ]"
# 3 groups of 4. Must not be glued to other digits/letters, and must not be
# part of the 16-digit VID printed on the card.
_NUMBER_RE = re.compile(
    rf"(?<![0-9A-Za-z])(?<![0-9]\s)"
    rf"({_NUM_CH}{{4}})[\s\-]?({_NUM_CH}{{4}})[\s\-]?({_NUM_CH}{{4}})"
    rf"(?![\s\-]?[0-9A-Za-z])"
)

_DATE = r"(\d{1,2})\s*[/\-\.]\s*(\d{1,2})\s*[/\-\.]\s*(\d{4})"
_DOB_LABEL = r"(?:DOB|D\.O\.B\.?|Date\s*of\s*Birth|जन्म\s*तिथि|जन्म\s*तारीख)\s*[:\-]?\s*"
_YOB_LABEL = r"(?:Year\s*of\s*Birth|YoB|जन्म\s*वर्ष)\s*[:\-]?\s*"

_NAME_STOP = {
    "government", "govt", "india", "bharat", "sarkar", "unique", "identification",
    "authority", "uidai", "aadhaar", "aadhar", "card", "enrolment", "enrollment",
    "dob", "birth", "year", "date", "male", "female", "transgender", "address",
    "vid", "download", "downloaded", "issue", "issued", "help", "www", "gov",
    "mera", "pehchaan", "proof", "identity", "citizenship", "verify", "online",
    "mobile", "no", "of", "the", "to", "son", "daughter", "wife", "father",
}
_NAME_RE = re.compile(r"^[A-Za-z][A-Za-z .'\-]{2,39}$")


def _find_number(lines: List[str]) -> str:
    found: List[str] = []
    for line in lines:
        for m in _NUMBER_RE.finditer(line):
            raw = "".join(m.groups())
            if sum(ch.isdigit() for ch in raw) < 8:   # avoid words like "SOLO BOSS"
                continue
            found.append(raw.translate(_OCR_FIX))
    if not found:
        return ""
    ranked = [n for n, _ in Counter(found).most_common()]
    for n in ranked:                                   # 1st choice: checksum passes
        if is_valid_aadhaar(n):
            return format_aadhaar(n)
    for n in ranked:                                   # 2nd: at least a legal 1st digit
        if n[0] not in "01":
            return format_aadhaar(n)
    return format_aadhaar(ranked[0])


def _valid_date(d: int, m: int, y: int) -> bool:
    try:
        datetime.date(y, m, d)
    except ValueError:
        return False
    return 1900 <= y <= datetime.date.today().year


def _find_dob(text: str) -> str:
    for m in re.finditer(_DOB_LABEL + _DATE, text, re.IGNORECASE):
        d, mo, y = (int(g) for g in m.groups())
        if _valid_date(d, mo, y):
            return f"{d:02d}/{mo:02d}/{y}"
    m = re.search(_YOB_LABEL + r"(\d{4})\b", text, re.IGNORECASE)
    if m:
        return m.group(1)
    for m in re.finditer(_DATE, text):
        d, mo, y = (int(g) for g in m.groups())
        if _valid_date(d, mo, y):
            return f"{d:02d}/{mo:02d}/{y}"
    return ""


def _find_gender(text: str) -> str:
    if re.search(r"\bfemale\b|महिला", text, re.IGNORECASE):
        return "FEMALE"
    if re.search(r"\bmale\b|पुरुष", text, re.IGNORECASE):
        return "MALE"
    if re.search(r"\btransgender\b|उभयलिंगी", text, re.IGNORECASE):
        return "TRANSGENDER"
    return ""


def _looks_like_name(line: str) -> bool:
    cand = line.strip()
    if not _NAME_RE.match(cand):
        return False
    tokens = [t.strip(".'-").lower() for t in cand.split()]
    if not tokens or len(tokens) > 5:
        return False
    if any(t in _NAME_STOP for t in tokens):
        return False
    return any(len(t) >= 2 for t in tokens)


def _find_name(lines: List[str]) -> str:
    anchor = None
    for i, line in enumerate(lines):
        if (re.search(_DOB_LABEL, line, re.IGNORECASE)
                or re.search(_YOB_LABEL, line, re.IGNORECASE)
                or re.search(_DATE, line)):
            anchor = i
            break
    if anchor is not None:                      # the name sits just above the DOB line
        for i in range(anchor - 1, max(-1, anchor - 4), -1):
            if _looks_like_name(lines[i]):
                return " ".join(lines[i].split())
    for line in lines:
        if _looks_like_name(line):
            return " ".join(line.split())
    return ""


def parse_aadhaar_text(lines: List[str]) -> Dict[str, Any]:
    lines = [ln.strip() for ln in lines if ln and ln.strip()]
    text = "\n".join(lines)
    return {
        "aadhaar_number": _find_number(lines),
        "name": _find_name(lines),
        "dob": _find_dob(text),
        "gender": _find_gender(text),
    }


# --------------------------------------------------------------------------
# Image preparation (EXIF rotation + downscale => faster, safer OCR)
# --------------------------------------------------------------------------
def prepare_image(src: str, out_dir: str, max_side: int = 1600, rotate_deg: int = 0) -> str:
    """Return the path of an upright JPEG no larger than `max_side` px.

    Phone photos are often 12 MP+ and carry an EXIF rotation flag that
    BitmapFactory ignores. Fixing both here avoids out-of-memory crashes and
    sideways OCR. Falls back to the original file if Pillow is unavailable.
    """
    try:
        from PIL import Image, ImageOps
    except ImportError:
        return src
    os.makedirs(out_dir, exist_ok=True)
    resample = getattr(Image, "Resampling", Image).LANCZOS
    with Image.open(src) as opened:
        im = ImageOps.exif_transpose(opened)
        if rotate_deg:
            im = im.rotate(rotate_deg, expand=True)
        im = im.convert("RGB")
        im.thumbnail((max_side, max_side), resample)
        out = os.path.join(out_dir, f"scan_{datetime.datetime.now():%Y%m%d_%H%M%S_%f}.jpg")
        im.save(out, "JPEG", quality=90)
    return out


# --------------------------------------------------------------------------
# Excel export
# --------------------------------------------------------------------------
def export_records(records: List[Dict[str, str]], path: str, mask: bool = False) -> str:
    import openpyxl  # lazy: keeps app start-up fast
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Aadhaar Cards"
    headers = ["S.No.", "Aadhaar Number", "Full Name", "Date of Birth",
               "Gender", "Checksum", "Scanned At"]
    ws.append(headers)

    for i, rec in enumerate(records, 1):
        number = rec.get("aadhaar", "")
        if not number:
            status = ""
        else:
            status = "Valid" if is_valid_aadhaar(number) else "Check digits"
        ws.append([i, mask_aadhaar(number) if mask else number, rec.get("name", ""),
                   rec.get("dob", ""), rec.get("gender", ""), status,
                   rec.get("scanned_at", "")])

    side = Side(border_style="thin", color="CBD5E1")
    border = Border(top=side, bottom=side, left=side, right=side)
    center = Alignment(horizontal="center", vertical="center")
    head_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    head_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    for cell in ws[1]:
        cell.fill, cell.font, cell.alignment, cell.border = head_fill, head_font, center, border

    centered_cols = {1, 2, 4, 5, 6, 7}
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            if cell.column in centered_cols:
                cell.alignment = center
            if cell.column in (2, 4):
                cell.number_format = "@"          # keep as text

    for idx, col in enumerate(ws.columns, 1):
        longest = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(idx)].width = min(max(longest + 4, 12), 40)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)
    return path
